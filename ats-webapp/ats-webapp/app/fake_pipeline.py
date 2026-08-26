"""
BACKEND FINTO (mock) — sostituisce temporaneamente l'avvio reale dei
container Docker (Sezione "Flusso caricamento -> elaborazione ->
archiviazione" dello schema). Rispetta lo stesso contratto di entita'
(Elaborazione, LogElaborazione, Prescrizione, DatiOcr, Difformita) che
usera' l'integrazione reale.

I file (PDF/PNG/JSON/Excel) vengono comunque scritti per davvero, ma
in una cartella locale ./sharepoint_finto/ invece che sul vero
SharePoint montato: cosi' il principio "SharePoint e' il filesystem,
il DB salva solo percorsi relativi" resta vero anche in sviluppo, senza
bisogno di avere OneDrive sincronizzato sulla macchina.

Quando il backend reale sara' pronto, queste tre funzioni vanno
sostituite dalle chiamate reali ai container Docker (che scriveranno
sul percorso SharePoint vero, letto da Configurazione.sharepoint_base_path).
"""
import json
import os
import random
import time
from datetime import datetime

from openpyxl import Workbook
from sqlalchemy.orm import Session, joinedload

from app.database import SessionLocal
from app.models import (
    LottoMensile, StatoLotto, Elaborazione, FaseElaborazione, StatoElaborazione,
    LogElaborazione, LivelloLog, Prescrizione, StatoBarcode, DatiOcr,
    Difformita, StatoDifformita,
)
from app.generatore_finto import genera_barcode, genera_dati_ocr_finti, CODICI_DIFFORMITA

FAKE_SP_ROOT = "sharepoint_finto"

# Cartella "fissa" (uguale per tutti i lotti, non ha un sottopercorso
# per lotto come sp_archivio_path) in cui viene salvato l'unico Excel
# di output di ogni lotto, cosi' com'e' descritto dal flusso reale:
# CANNABIS/Archivio/Elaborazioni Recenti. "CANNABIS" non va ripetuto
# qui perche' e' gia' la cartella a cui punta
# Configurazione.sharepoint_base_path (FAKE_SP_ROOT la simula).
PERCORSO_CARTELLA_OUTPUT_RECENTI = "ARCHIVIO/ELABORAZIONI RECENTI"

# Probabilita' (0-1) di simulare un'eccezione durante una fase, solo a
# scopo dimostrativo: nel backend reale l'errore sara' quello vero
# sollevato dal container Docker corrispondente.
PROBABILITA_ECCEZIONE_FINTA = 0.08


def _percorso_assoluto(percorso_relativo: str) -> str:
    path = os.path.join(FAKE_SP_ROOT, percorso_relativo)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    return path


def _scrivi_file_finto(percorso_relativo: str, contenuto: str = "") -> None:
    path = _percorso_assoluto(percorso_relativo)
    with open(path, "w", encoding="utf-8") as f:
        f.write(contenuto)


def genera_pdf_finto(righe: list) -> bytes:
    """
    Costruisce un PDF minimo ma realmente valido (una pagina, testo
    semplice), senza dipendenze esterne: serve solo a far si' che il
    file scritto da questo modulo (o da seed_admin.py, che riusa questa
    stessa funzione) sia davvero apribile in un visualizzatore PDF (link
    "Apri prescrizione" nel dettaglio lotto), invece di un file di testo
    con estensione .pdf. Non e' un rendering del documento originale,
    solo un segnaposto leggibile. Pubblica (niente underscore) perche'
    usata anche fuori da questo modulo.
    """
    def _escape(s: str) -> str:
        return s.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")

    y = 780
    comandi = []
    for riga in righe:
        comandi.append(f"BT /F1 13 Tf 50 {y} Td ({_escape(riga)}) Tj ET")
        y -= 22
    stream_content = "\n".join(comandi).encode("latin-1", errors="replace")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream_content)).encode() + b" >>\nstream\n"
        + stream_content + b"\nendstream",
    ]

    buf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, body in enumerate(objects, start=1):
        offsets.append(len(buf))
        buf += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"

    xref_offset = len(buf)
    buf += f"xref\n0 {len(objects) + 1}\n".encode()
    buf += b"0000000000 65535 f \n"
    for off in offsets[1:]:
        buf += f"{off:010d} 00000 n \n".encode()
    buf += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF"
    ).encode()

    return bytes(buf)


def scrivi_pdf_finto(percorso_relativo: str, righe: list) -> None:
    path = _percorso_assoluto(percorso_relativo)
    with open(path, "wb") as f:
        f.write(genera_pdf_finto(righe))


def _log(db: Session, elaborazione: Elaborazione, messaggio: str, livello: LivelloLog = LivelloLog.info) -> None:
    db.add(LogElaborazione(elaborazione_id=elaborazione.id, messaggio=messaggio, livello=livello))
    db.commit()


# Titolo comprensibile e possibili cause per ciascuna fase: usati per
# rendere il messaggio mostrato all'operatore specifico e plausibile,
# invece del generico "container terminato con errore" sempre uguale.
# Nel backend reale, "causa" verra' sostituita dal messaggio vero
# sollevato dal container Docker corrispondente (qui solo simulato).
TITOLO_PER_FASE = {
    "preprocessing": "Il preprocessing del PDF si è interrotto",
    "elaborazione OCR": "L'estrazione automatica dei dati (OCR) si è interrotta",
    "analisi difformita": "L'analisi delle non conformità si è interrotta",
}

CAUSE_ERRORE_PER_FASE = {
    "preprocessing": [
        "Il PDF caricato risulta danneggiato e non è stato possibile aprirlo.",
        "Una o più pagine del PDF non contengono un barcode leggibile: lo split delle prescrizioni si è interrotto.",
        "Il documento sembra una scansione di bassa qualità: la correzione automatica dell'inclinazione (deskew) non è riuscita a completarsi.",
        "Il file caricato non è nel formato atteso (verificare che non sia protetto da password o scannerizzato come immagine unica).",
    ],
    "elaborazione OCR": [
        "Il servizio di riconoscimento automatico (OCR) non ha risposto entro il tempo massimo previsto.",
        "Il servizio OCR non è raggiungibile: il container di elaborazione potrebbe non essere attivo.",
        "La qualità di una o più immagini era troppo bassa per un'estrazione affidabile dei dati.",
        "Memoria insufficiente sul server durante l'elaborazione di un documento particolarmente pesante.",
    ],
    "analisi difformita": [
        "Non è stato possibile leggere il file Excel regionale caricato: verificare che non sia danneggiato.",
        "Il formato dell'Excel regionale non corrisponde a quello atteso (colonne mancanti o in ordine diverso).",
        "Il confronto tra i dati estratti e le regole di conformità si è interrotto per un errore interno.",
    ],
}


def _messaggio_errore_operatore(fase_label: str, causa: str) -> str:
    titolo = TITOLO_PER_FASE.get(fase_label, "L'elaborazione si è interrotta")
    return (
        f"{titolo}.\n{causa}\n\n"
        f"Cosa fare: segnala al team tecnico il nome del lotto e l'orario riportati sopra."
    )


def _messaggio_errore_imprevisto(fase_label: str) -> str:
    """Usato quando l'eccezione NON e' quella simulata da _forse_eccezione ma un vero errore
    imprevisto nel codice: non mostriamo il traceback grezzo all'operatore, solo un messaggio
    onesto ma comprensibile. Il dettaglio tecnico resta nel registro dell'elaborazione."""
    titolo = TITOLO_PER_FASE.get(fase_label, "L'elaborazione si è interrotta")
    return (
        f"{titolo} per un problema tecnico imprevisto, non legato ai dati caricati.\n\n"
        f"Cosa fare: segnala al team tecnico il nome del lotto e l'orario riportati sopra; "
        f"il dettaglio tecnico è stato salvato nel registro dell'elaborazione."
    )


def _forse_eccezione(db: Session, lotto: LottoMensile, elaborazione: Elaborazione, fase_label: str) -> bool:
    """Ritorna True se ha simulato un'eccezione (e ha gia' salvato tutto)."""
    if random.random() >= PROBABILITA_ECCEZIONE_FINTA:
        return False
    causa = random.choice(CAUSE_ERRORE_PER_FASE.get(fase_label, ["Si è verificato un problema imprevisto."]))
    _log(db, elaborazione, f"[{fase_label}] {causa}", LivelloLog.error)
    elaborazione.stato = StatoElaborazione.errore
    elaborazione.exit_code = 1
    elaborazione.finished_at = datetime.utcnow()
    lotto.stato = StatoLotto.eccezione
    lotto.note = _messaggio_errore_operatore(fase_label, causa)
    lotto.updated_at = datetime.utcnow()
    db.commit()
    return True


def avvia_preprocessing_fake(lotto_id) -> None:
    """Fase 1: split PDF, lettura barcode, deskew. Crea le Prescrizioni."""
    db: Session = SessionLocal()
    try:
        lotto = db.query(LottoMensile).filter(LottoMensile.id == lotto_id).first()
        if lotto is None:
            return
        elaborazione = None  # garantisce che sia definita anche se l'eccezione arriva prima della creazione

        elaborazione = Elaborazione(
            lotto_id=lotto.id, fase=FaseElaborazione.preprocessing,
            stato=StatoElaborazione.in_corso,
            comando_docker=f"docker compose run preprocessing --lotto {lotto.nome}",
            started_at=datetime.utcnow(),
        )
        db.add(elaborazione)
        db.commit()

        lotto.stato = StatoLotto.preprocessing
        db.commit()

        _log(db, elaborazione, f"Avvio preprocessing per il lotto '{lotto.nome}'")
        time.sleep(1)
        _log(db, elaborazione, "Split del PDF combinato in singole prescrizioni")
        time.sleep(1.5)

        if _forse_eccezione(db, lotto, elaborazione, "preprocessing"):
            return

        numero_prescrizioni = random.randint(6, 24)
        n_letti = 0
        for i in range(numero_prescrizioni):
            barcode_letto = random.random() > 0.12  # ~88% barcode leggibili
            barcode = genera_barcode() if barcode_letto else None
            presc = Prescrizione(
                lotto_id=lotto.id,
                barcode=barcode,
                stato_barcode=StatoBarcode.letto if barcode_letto else StatoBarcode.undefined,
                sp_pdf_path=f"{lotto.sp_output_path}/CARTELLE FARMACIE/{barcode or f'undefined_{i}'}.pdf",
                sp_png_path=f"{lotto.sp_output_path}/CARTELLE FARMACIE/{barcode or f'undefined_{i}'}.png",
            )
            db.add(presc)
            scrivi_pdf_finto(presc.sp_pdf_path, [
                "PRESCRIZIONE - DOCUMENTO SIMULATO (ambiente di sviluppo)",
                f"Lotto: {lotto.nome}",
                f"Barcode: {barcode or 'non ancora letto'}",
                f"Posizione nel PDF caricato originale: #{i + 1}",
                "",
                "File generato dal backend finto (app/fake_pipeline.py) al posto",
                "del vero split del PDF (fase 1 Docker), solo per sviluppo/demo.",
            ])
            if barcode_letto:
                n_letti += 1
        db.commit()

        lotto.n_pdf_caricati = 1
        lotto.n_prescrizioni_totali = numero_prescrizioni
        lotto.n_barcode_letti = n_letti
        lotto.n_barcode_undefined = numero_prescrizioni - n_letti
        lotto.updated_at = datetime.utcnow()

        elaborazione.stato = StatoElaborazione.completata
        elaborazione.finished_at = datetime.utcnow()
        elaborazione.exit_code = 0
        elaborazione.n_processati = numero_prescrizioni
        elaborazione.n_errori = 0

        lotto.stato = StatoLotto.revisione_barcode
        _log(db, elaborazione, f"Preprocessing completato: {numero_prescrizioni} prescrizioni, {n_letti} barcode letti, {numero_prescrizioni - n_letti} da rivedere")
        db.commit()

    except Exception as exc:  # pragma: no cover - solo per il mock
        lotto = db.query(LottoMensile).filter(LottoMensile.id == lotto_id).first()
        if lotto:
            lotto.stato = StatoLotto.eccezione
            lotto.note = _messaggio_errore_imprevisto("preprocessing")
            lotto.updated_at = datetime.utcnow()
            if elaborazione is not None:
                _log(db, elaborazione, f"Eccezione non gestita: {exc}", LivelloLog.error)
            db.commit()
    finally:
        db.close()


def avvia_ocr_fake(lotto_id) -> None:
    """Fase 2: estrazione VLLM dei 24 campi per ogni prescrizione."""
    db: Session = SessionLocal()
    try:
        lotto = db.query(LottoMensile).filter(LottoMensile.id == lotto_id).first()
        if lotto is None:
            return
        elaborazione = None  # garantisce che sia definita anche se l'eccezione arriva prima della creazione

        elaborazione = Elaborazione(
            lotto_id=lotto.id, fase=FaseElaborazione.vllm,
            stato=StatoElaborazione.in_corso,
            comando_docker=f"docker compose run ocr-vllm --lotto {lotto.nome}",
            started_at=datetime.utcnow(),
        )
        db.add(elaborazione)
        lotto.stato = StatoLotto.elaborazione_ocr
        db.commit()

        _log(db, elaborazione, "Avvio estrazione OCR (Qwen2.5-VL) sulle prescrizioni")
        time.sleep(1)

        prescrizioni = db.query(Prescrizione).filter(Prescrizione.lotto_id == lotto.id).all()
        n_match = 0
        somma_score = 0.0
        for presc in prescrizioni:
            time.sleep(0.15)
            if _forse_eccezione(db, lotto, elaborazione, "elaborazione OCR"):
                return

            dati = genera_dati_ocr_finti()
            score = round(random.uniform(72, 99), 2)
            somma_score += score

            dati_ocr = DatiOcr(
                prescrizione_id=presc.id,
                json_vllm_raw={k: (v.isoformat() if hasattr(v, "isoformat") else v) for k, v in dati.items()},
                extracted_at=datetime.utcnow(),
                **dati,
            )
            db.add(dati_ocr)

            presc.score_ocr = score
            presc.n_campi_compilati = random.randint(20, 24)
            presc.barcode_in_excel = random.random() > 0.1
            if presc.barcode_in_excel:
                presc.riga_excel = random.randint(2, 300)
                n_match += 1
            presc.sp_json_path = presc.sp_pdf_path.replace(".pdf", ".json") if presc.sp_pdf_path else None
            if presc.sp_json_path:
                _scrivi_file_finto(presc.sp_json_path, json.dumps(dati_ocr.json_vllm_raw, ensure_ascii=False, indent=2))

        db.commit()

        lotto.n_match_excel = n_match
        lotto.score_ocr_medio = round(somma_score / len(prescrizioni), 2) if prescrizioni else None
        lotto.updated_at = datetime.utcnow()

        elaborazione.stato = StatoElaborazione.completata
        elaborazione.finished_at = datetime.utcnow()
        elaborazione.exit_code = 0
        elaborazione.n_processati = len(prescrizioni)

        lotto.stato = StatoLotto.revisione_qualita
        _log(db, elaborazione, f"OCR completato su {len(prescrizioni)} prescrizioni, score medio {lotto.score_ocr_medio}")
        db.commit()

    except Exception as exc:  # pragma: no cover
        lotto = db.query(LottoMensile).filter(LottoMensile.id == lotto_id).first()
        if lotto:
            lotto.stato = StatoLotto.eccezione
            lotto.note = _messaggio_errore_imprevisto("elaborazione OCR")
            lotto.updated_at = datetime.utcnow()
            if elaborazione is not None:
                _log(db, elaborazione, f"Eccezione non gestita: {exc}", LivelloLog.error)
            db.commit()
    finally:
        db.close()


# ============================================================
# Foglio di output finale — un Excel per lotto, salvato in
# PERCORSO_CARTELLA_OUTPUT_RECENTI (fase finale del flusso, quando il
# lotto passa a "completato"). Struttura e nomi colonna sono quelli
# reali del sistema regionale + i campi estratti dall'OCR.
# ============================================================

COLONNE_OUTPUT_EXCEL = [
    "Data_Elaborazione", "DATA_CONTABILE", "FARMACIA_ID", "FARMACIA_KEY", "DATA_SPEDIZIONE", "NUM_RICETTA",
    "BARCODE", "MEDICO_ID", "ASSISTITO_ID", "COD_FISCALE", "FARMACO_ID", "QTA",
    "PREZZO", "RIC_SISS_ID", "FLAG_SISS", "LORDO_PRESC", "NUM_DDD", "LORDO_TOTALE",
    "TICKET_REGIST", "IMPORTO_RIC_EXTRASCONTO", "TRATT_SCONTO", "FLAG_EXTRASCONTO", "FLAG_PRESC_SUGG", "TIPO_RICETTA",
    "CANALE_ID", "FLAG_TESTATA", "ASL_ID_ASSTO", "FULL_DISTRETTO_ID_ASSTO", "ASL_ID", "FULL_DISTRETTO_ID",
    "MACRO_ESENZIONE_ID", "CODICE_ESENZIONE_ID", "CODICE_ESENZIONE_KEY", "SESSO", "ETA_ANNI", "FULL_DISTRETTO_ID_FARMA",
    "ASL_ID_FARMA", "FLAG_PROTESICO", "ATC_ID", "ENTE_APP_AMM_ID", "SPECIALISTA_ID", "STRUTTURA_ID",
    "cognome_nome_assistito", "barcode", "codice_fiscale", "codice_esenzione", "codice_atc", "testo_prescrizione",
    "metodo_estrattivo_olio", "forma_farmaceutica", "data_etichetta_preparazione", "data_prescrizione", "data_invio", "timbro_medico",
    "firma_medico", "etichetta_avvertenze", "etichetta_data_scadenza", "etichetta_nome_cognome_medico", "etichetta_nome_cognome_paziente", "etichetta_prezzo_sost",
    "etichetta_prezzo_on", "etichetta_prezzo_rec", "etichetta_prezzo_iva", "etichetta_prezzo_tot", "totale_prescrizione", "etichetta_THC",
    "nome_farmacia", "LINK", "R02", "R03", "R04", "R05",
    "R07", "R05A", "R06", "R08", "R09", "R10",
    "R10A", "R01", "R11", "R12", "R13", "R16",
    "R19", "R14", "R18", "ETICHETTA MANCANTE", "ESITO",
]

# Colonne R** -> codice Difformita corrispondente (si ottiene togliendo
# la "R" iniziale: R05A -> "05A"). Solo i codici gia' generati da
# generatore_finto.CODICI_DIFFORMITA compariranno mai valorizzati nella
# modalita' finta; le altre colonne R restano vuote finche' il set di
# non conformita' simulate non viene esteso alle 19 reali.
MAPPA_COLONNE_DIFFORMITA = {
    colonna: colonna[1:] for colonna in COLONNE_OUTPUT_EXCEL if colonna.startswith("R") and colonna[1:2].isdigit()
}

# Le poche colonne mostrate nell'anteprima minimale nel dettaglio lotto
# (l'Excel completo ha 89 colonne, troppe per una tabella a schermo).
COLONNE_ANTEPRIMA_OUTPUT = [
    "BARCODE", "cognome_nome_assistito", "nome_farmacia", "data_prescrizione", "totale_prescrizione", "ESITO",
]


def nome_file_output_excel(lotto: LottoMensile) -> str:
    nome_normalizzato = "_".join(lotto.nome.split())  # "LUGLIO 2025" -> "LUGLIO_2025"
    return f"OUTPUT_FLUSSO_CANNABIS_{nome_normalizzato}.xlsx"


def salva_workbook_output(percorso_relativo: str, wb: Workbook) -> None:
    """Pubblica: usata sia da genera_excel_output_fake sia da seed_admin.py, che ha gia' una propria sessione aperta."""
    wb.save(_percorso_assoluto(percorso_relativo))


def costruisci_riga_output(p: Prescrizione, lotto: LottoMensile) -> dict:
    """
    Una riga del foglio di output per una prescrizione. Unica fonte
    usata sia per scrivere l'Excel sia per l'anteprima nel dettaglio
    lotto. Le colonne del sistema regionale senza un campo
    corrispondente nella modalita' finta (FARMACIA_ID, MEDICO_ID,
    ASL_ID, ...) restano vuote: nel backend reale arriverebbero dal
    match con l'Excel regionale caricato (Prescrizione.riga_excel),
    che qui non viene letto riga per riga.
    """
    d = p.dati_ocr
    codici_difformita_attivi = {df.codice for df in p.difformita if df.stato != StatoDifformita.esclusa}

    riga = {
        "Data_Elaborazione": lotto.completato_at.strftime("%d/%m/%Y") if lotto.completato_at else datetime.utcnow().strftime("%d/%m/%Y"),
        "BARCODE": p.barcode or "",
        "COD_FISCALE": d.codice_fiscale if d else "",
        "LINK": f"/lotti/{p.lotto_id}/prescrizioni/{p.id}/pdf",
    }

    if d:
        riga.update({
            "cognome_nome_assistito": d.cognome_nome_assistito or "",
            "barcode": p.barcode or "",
            "codice_fiscale": d.codice_fiscale or "",
            "codice_esenzione": d.codice_esenzione or "",
            "codice_atc": d.codice_atc or "",
            "testo_prescrizione": d.testo_prescrizione or "",
            "metodo_estrattivo_olio": d.metodo_estrattivo_olio or "",
            "forma_farmaceutica": d.forma_farmaceutica or "",
            "data_etichetta_preparazione": d.data_etichetta_preparazione.strftime("%d/%m/%Y") if d.data_etichetta_preparazione else "",
            "data_prescrizione": d.data_prescrizione.strftime("%d/%m/%Y") if d.data_prescrizione else "",
            "data_invio": d.data_invio.strftime("%d/%m/%Y") if d.data_invio else "",
            "timbro_medico": "SI" if d.timbro_medico else ("NO" if d.timbro_medico is False else ""),
            "firma_medico": "SI" if d.firma_medico else ("NO" if d.firma_medico is False else ""),
            "etichetta_avvertenze": d.etichetta_avvertenze or "",
            "etichetta_data_scadenza": d.etichetta_data_scadenza.strftime("%d/%m/%Y") if d.etichetta_data_scadenza else "",
            "etichetta_nome_cognome_medico": d.etichetta_nome_cognome_medico or "",
            "etichetta_nome_cognome_paziente": d.etichetta_nome_cognome_paziente or "",
            "etichetta_prezzo_sost": float(d.etichetta_prezzo_sost) if d.etichetta_prezzo_sost is not None else "",
            "etichetta_prezzo_on": float(d.etichetta_prezzo_on) if d.etichetta_prezzo_on is not None else "",
            "etichetta_prezzo_rec": float(d.etichetta_prezzo_rec) if d.etichetta_prezzo_rec is not None else "",
            "etichetta_prezzo_iva": float(d.etichetta_prezzo_iva) if d.etichetta_prezzo_iva is not None else "",
            "etichetta_prezzo_tot": float(d.etichetta_prezzo_tot) if d.etichetta_prezzo_tot is not None else "",
            "totale_prescrizione": float(d.totale_prescrizione) if d.totale_prescrizione is not None else "",
            "etichetta_THC": d.etichetta_thc or "",
            "nome_farmacia": d.nome_farmacia or "",
        })

    for colonna_r, codice in MAPPA_COLONNE_DIFFORMITA.items():
        riga[colonna_r] = "X" if codice in codici_difformita_attivi else ""

    riga["ETICHETTA MANCANTE"] = "SI" if (d is None or d.etichetta_prezzo_tot is None) else "NO"
    riga["ESITO"] = "DIFFORME" if codici_difformita_attivi else "OK"

    return riga


def costruisci_workbook_output(lotto: LottoMensile, prescrizioni: list) -> Workbook:
    """Pura: dato un lotto e le sue prescrizioni gia' caricate, costruisce il Workbook. Nessuna sessione DB."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Output"
    ws.append(COLONNE_OUTPUT_EXCEL)
    for cella in ws[1]:
        cella.font = cella.font.copy(bold=True)
    ws.freeze_panes = "A2"

    for p in prescrizioni:
        riga = costruisci_riga_output(p, lotto)
        ws.append([riga.get(colonna, "") for colonna in COLONNE_OUTPUT_EXCEL])

    for indice, colonna in enumerate(COLONNE_OUTPUT_EXCEL, start=1):
        ws.column_dimensions[ws.cell(row=1, column=indice).column_letter].width = max(12, min(28, len(colonna) + 2))

    return wb


def genera_excel_output_fake(lotto_id) -> None:
    """
    Genera e salva il foglio di output finale per il lotto, in
    PERCORSO_CARTELLA_OUTPUT_RECENTI. Chiamata sincrona da
    completa_lotto() (e' una singola scrittura, non serve
    BackgroundTasks). Apre una propria sessione perche' e' invocata
    anche fuori dal ciclo delle altre fasi (non da un job in coda).
    """
    db: Session = SessionLocal()
    try:
        lotto = (
            db.query(LottoMensile)
            .options(joinedload(LottoMensile.prescrizioni).joinedload(Prescrizione.dati_ocr))
            .options(joinedload(LottoMensile.prescrizioni).joinedload(Prescrizione.difformita))
            .filter(LottoMensile.id == lotto_id)
            .first()
        )
        if lotto is None:
            return

        wb = costruisci_workbook_output(lotto, lotto.prescrizioni)
        nome_file = nome_file_output_excel(lotto)
        percorso_relativo = f"{PERCORSO_CARTELLA_OUTPUT_RECENTI}/{nome_file}"
        salva_workbook_output(percorso_relativo, wb)

        lotto.excel_output_filename = nome_file
        db.commit()
    finally:
        db.close()


def avvia_difformita_fake(lotto_id) -> None:
    """Fase 4: analisi di 19 tipologie di non conformita (qui simulate)."""
    db: Session = SessionLocal()
    try:
        lotto = db.query(LottoMensile).filter(LottoMensile.id == lotto_id).first()
        if lotto is None:
            return
        elaborazione = None  # garantisce che sia definita anche se l'eccezione arriva prima della creazione

        elaborazione = Elaborazione(
            lotto_id=lotto.id, fase=FaseElaborazione.difformita,
            stato=StatoElaborazione.in_corso,
            comando_docker=f"docker compose run difformita --lotto {lotto.nome}",
            started_at=datetime.utcnow(),
        )
        db.add(elaborazione)
        lotto.stato = StatoLotto.analisi_difformita
        db.commit()

        _log(db, elaborazione, "Avvio analisi delle non conformita regolamentari")
        time.sleep(1.5)

        if _forse_eccezione(db, lotto, elaborazione, "analisi difformita"):
            return

        prescrizioni = db.query(Prescrizione).filter(Prescrizione.lotto_id == lotto.id).all()
        n_con_difformita = 0
        n_difformita_totali = 0
        for presc in prescrizioni:
            if random.random() < 0.25:
                n_di_questa = random.randint(1, 2)
                for _ in range(n_di_questa):
                    codice, descrizione = random.choice(CODICI_DIFFORMITA)
                    db.add(Difformita(
                        prescrizione_id=presc.id, codice=codice, descrizione=descrizione,
                        stato=StatoDifformita.rilevata,
                    ))
                n_con_difformita += 1
                n_difformita_totali += n_di_questa
        db.commit()

        lotto.n_difformita_totali = n_difformita_totali
        lotto.n_prescrizioni_con_difformita = n_con_difformita
        lotto.updated_at = datetime.utcnow()

        elaborazione.stato = StatoElaborazione.completata
        elaborazione.finished_at = datetime.utcnow()
        elaborazione.exit_code = 0
        elaborazione.n_processati = len(prescrizioni)

        lotto.stato = StatoLotto.revisione_difformita
        _log(db, elaborazione, f"Analisi completata: {n_difformita_totali} difformita su {n_con_difformita} prescrizioni")
        db.commit()

    except Exception as exc:  # pragma: no cover
        lotto = db.query(LottoMensile).filter(LottoMensile.id == lotto_id).first()
        if lotto:
            lotto.stato = StatoLotto.eccezione
            lotto.note = _messaggio_errore_imprevisto("analisi difformita")
            lotto.updated_at = datetime.utcnow()
            if elaborazione is not None:
                _log(db, elaborazione, f"Eccezione non gestita: {exc}", LivelloLog.error)
            db.commit()
    finally:
        db.close()
